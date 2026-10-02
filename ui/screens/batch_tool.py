"""
BATCH_TOOL.PY — Пакетная обработка чертежей (PDF, КОМПАС .cdw, КОМПАС .m3d/.a3d)

Режимы:
  1. Список файлов — выбрать чертежи вручную.
  2. Режим папки:
       Папка 1 — исходные чертежи и модели (.cdw, .pdf, .m3d/.a3d), без изменения файлов
       Папка 2 — Excel-таблица и протокол обработки

.m3d/.a3d даёт только паспортные данные (обозначение, материал, масса) —
геометрия в этом формате закрыта (C3D), поэтому заготовка определяется
только если деталь уже есть в базе эталонов по обозначению. Полный расчёт
по 3D-геометрии — через вкладку «3D инструмент» (STEP) или живой КОМПАС.
"""

import os
import re
import threading
import datetime
import customtkinter as ctk
from tkinter import filedialog, messagebox

from core.engine import (
    extract_from_pdf,
    normalize_mark,
    normalize_gost,
    ai_ask,
)
# Расширение сортаментов: шестигранник, полоса, уголок.
# Импорт обязателен ДО первого расчёта — модуль дополняет справочники движка.
# Сам engine.py при этом не изменён.
from core.stock_ext import calculate_stock_line as calculate_missing_stock_line
from core import cdw_bridge
from core import drawing_db
from core.drawing_db import build_etalon_examples_block, lookup_etalon
from core.gost_reader import get_gost_context_for_query
from core import logger
from core import activity_log


# ============================================================
# Вспомогательные функции
# ============================================================

def extract_drawing_no(text: str) -> str:
    if not text:
        return ""
    m = re.search(r"\b(\d{2,6}\.\d{2}\.\d{2,3}\.\d{3,})\b", text)
    return m.group(1) if m else ""


def _drawing_sort_key(drawing_no: str):
    """
    Ключ натуральной сортировки по номеру чертежа.
    '290426.00.00.071' → (290426, 0, 0, 71). Пустые/без номера — в конец.
    """
    if not drawing_no:
        return (1, ())
    nums = tuple(int(x) for x in re.findall(r"\d+", drawing_no))
    return (0, nums)


def extract_clean_mass(pdf_text: str) -> str:
    """
    Масса чистовая — из поля 'Масса' штампа чертежа.
    Это одиночное число (целое или с запятой) рядом с датами и литерой И.
    """
    lines = [ln.strip() for ln in pdf_text.splitlines()]
    mass_re = re.compile(r'^\d{1,5}(?:[,\.]\d{1,3})?$')
    date_re = re.compile(r'\d{2}\.\d{2}\.\d{4}')

    candidates = []
    for i, ln in enumerate(lines):
        if not mass_re.match(ln):
            continue
        nearby = lines[max(0, i - 5):i + 6]
        has_date = any(date_re.search(l) for l in nearby)
        has_lit  = any(l in ('И', 'И ') for l in nearby)
        if not has_date:
            continue
        val = ln.replace('.', ',')
        try:
            if float(val.replace(',', '.')) > 500:
                continue
        except ValueError:
            pass
        candidates.append((i, val, has_lit))

    lit_cands = [c for c in candidates if c[2]]
    if lit_cands:
        return lit_cands[0][1]
    if candidates:
        return candidates[0][1]
    return ""


def _fmt_mass_ru(v) -> str:
    """Число → русская строка без лишних нулей: 1.3→'1,3', 0.052→'0,052', 63.0→'63'."""
    if v is None or v == "":
        return ""
    try:
        s = f"{float(v):.3f}".rstrip("0").rstrip(".")
        return s.replace(".", ",")
    except (ValueError, TypeError):
        return str(v)


def extract_sortament_from_text(text: str) -> str:
    """
    Сортамент из штампа чертежа — поле Заготовка + Материал.
    Блок после литеры И: тип заготовки + марка материала.
    """
    if not text:
        return ""
    lines = [ln.strip() for ln in text.splitlines()]

    stock_re = re.compile(
        r"^(Круг|Лист|Труба|Квадрат|Шестигранник|Полоса|Уголок|Швеллер|Двутавр)",
        re.IGNORECASE,
    )
    material_re = re.compile(
        r"^(Сталь|Ст\d|40Х|38Х|30Х|12Х|09Г|В\s*\d|Д\d|АМ|Л\d|\d{2}Х)",
        re.IGNORECASE,
    )

    def norm_stock(s):
        return re.sub(
            r"^(Круг|Лист|Труба|Квадрат|Шестигранник|Полоса)(\d)",
            r"\1 \2", s, flags=re.IGNORECASE
        )

    stock_line = ""
    material_line = ""

    for i, ln in enumerate(lines):
        if ln in ("И", "И "):
            for j in range(i + 1, min(len(lines), i + 5)):
                raw = lines[j]
                c = norm_stock(raw)
                if stock_re.match(c) and not stock_line:
                    stock_line = c
                elif material_re.match(raw) and not material_line:
                    material_line = raw
            break

    if not stock_line and not material_line:
        for ln in lines:
            c = norm_stock(ln)
            if stock_re.match(c) and ln.lower() not in ("лист", "листов") and not stock_line:
                stock_line = c
            if material_re.match(ln) and not material_line:
                material_line = ln

    if stock_line and material_line:
        return f"{stock_line} / {material_line}"
    return stock_line or material_line or ""


def parse_result_line(result_line: str) -> tuple:
    """
    Разбирает строку результата на (сортамент_новый, размер, масса_заг).

    Формат строки в базе:
      'Круг 130 ГОСТ2590-2006;\nL=4680;\nМасса заготовки, кг: 490,00'

    Возвращает:
      new_sortament = 'Круг 130 ГОСТ2590-2006'
      stock_size    = 'L=4680'
      stock_mass    = '490,00'
    """
    if not result_line:
        return "", "", ""

    parts = [p.strip().rstrip(";") for p in result_line.split("\n")]
    new_sortament = parts[0] if len(parts) > 0 else ""
    stock_size    = parts[1] if len(parts) > 1 else ""
    massa_str     = parts[2] if len(parts) > 2 else ""

    # Извлекаем число из "Масса заготовки, кг: 490,00"
    m = re.search(r":\s*([\d,\.]+)", massa_str)
    stock_mass = m.group(1) if m else ""

    return new_sortament, stock_size, stock_mass


def _blank_result(path: str) -> dict:
    return {
        "file":               os.path.basename(path),
        "path":               path,
        "drawing_no":         "",
        "part_name":          "",
        "zagotovka_full":     "",   # Заготовка (сортамент/материал)
        "new_sortament":      "",   # Новый сортамент (ответ OTK.AI)
        "drawing_sortament":  "",   # Сортамент из чертежа (штамп)
        "stock_size":         "",   # Размер заготовки
        "stock_mass_kg":      "",   # Масса заготовки
        "clean_mass_kg":      "",   # Масса чистовая
        "material":           "",   # Материал из штампа
        "source":             "",
        "status":             "ОК",
        "error":              "",
    }


# ============================================================
# Обработка одного PDF
# ============================================================

def process_one_pdf(pdf_path: str) -> dict:
    name = os.path.basename(pdf_path)
    result = _blank_result(pdf_path)

    try:
        ex = extract_from_pdf(pdf_path, max_pages=3)
        result["_raw_text"] = getattr(ex, "pdf_text", "") or ""
        raw_name = ex.part_name or name.replace(".pdf", "")
        # Убираем префикс из имени: "23_14_02_004_Тяга" → "Тяга"
        m_clean = re.match(r'^[\d_]+([А-ЯЁа-яёA-Za-z].*)', raw_name)
        if m_clean:
            raw_name = m_clean.group(1)
        # Заменяем оставшиеся одиночные подчёркивания на пробелы
        raw_name = raw_name.replace("_", " ").strip()
        result["part_name"]         = raw_name
        result["drawing_no"]        = extract_drawing_no(ex.pdf_text)
        result["drawing_sortament"] = extract_sortament_from_text(ex.pdf_text)
        result["clean_mass_kg"]     = extract_clean_mass(ex.pdf_text)
        result["material"]          = " ".join(
            x for x in (ex.material_mark or "", ex.material_gost or "") if x
        )

        etalon = lookup_etalon(ex.pdf_text, ex.part_name)
        if etalon:
            result["drawing_no"] = etalon.get("drawing_no", result["drawing_no"])
            # Берём наименование из базы — оно правильное (без номера чертежа)
            db_name = etalon.get("part_name", "")
            if db_name:
                m_c = re.match(r'^[\d_]+([А-ЯЁа-яёA-Za-z].*)', db_name)
                result["part_name"] = m_c.group(1).replace("_", " ").strip() if m_c else db_name

            new_sortament, stock_size, stock_mass = parse_result_line(
                etalon.get("result_line", "")
            )
            result["new_sortament"]  = new_sortament
            result["stock_size"]     = stock_size
            result["stock_mass_kg"]  = stock_mass
            # Чистовая масса: приоритет — из базы эталонов, иначе из штампа
            db_clean = etalon.get("mass_clean_kg", None)
            if db_clean is not None:
                result["clean_mass_kg"] = _fmt_mass_ru(db_clean)
            result["zagotovka_full"] = new_sortament
            result["source"]         = "база"
            return result

        # Алгоритм
        mark = normalize_mark(ex.material_mark or "")
        gost = normalize_gost(ex.material_gost or "")
        length_mm = ex.length_mm

        if not mark:
            result["status"] = "⚠ нет марки"
            result["error"]  = "Марка материала не найдена"
            return result
        if not gost:
            result["status"] = "⚠ нет ГОСТ"
            result["error"]  = "ГОСТ материала не найден"
            return result
        if not length_mm or length_mm <= 0:
            result["status"] = "⚠ нет длины"
            result["error"]  = "Длина детали не определена"
            return result

        res = calculate_missing_stock_line(ex, length_mm, mark, gost)
        new_sortament, stock_size, stock_mass = parse_result_line(res.result_line)

        result["new_sortament"]  = new_sortament
        result["stock_size"]     = stock_size
        result["stock_mass_kg"]  = stock_mass
        result["zagotovka_full"] = new_sortament
        result["source"]         = "расчёт"

    except Exception as e:
        result["status"] = "❌ ошибка"
        result["error"]  = str(e)[:120]
        logger.error(f"Пакетная обработка: ошибка {name}: {e}")

    return result


# ============================================================
# Обработка одного .cdw (КОМПАС)
# ============================================================

_STOCK_WORDS = (
    "уголок", "швеллер", "труба", "круг", "квадрат",
    "полоса", "лист", "двутавр", "шестигранник", "балка",
)


def _sortament_from_library(card: dict) -> str:
    """
    Если деталь вставлена из библиотеки проката, сортамент лежит
    прямо в файле: 'Уголок неравнопол. ГОСТ 8510.frw'.
    """
    for ref in card.get("library_refs", []):
        base = os.path.splitext(ref)[0]
        if any(w in base.lower() for w in _STOCK_WORDS):
            return base.strip()
    return ""


def process_one_cdw(cdw_path: str) -> dict:
    """
    Чертёж КОМПАС читается структурно: обозначение, наименование,
    материал, плотность, масса чистовая берутся из файла точно.
    """
    name = os.path.basename(cdw_path)
    result = _blank_result(cdw_path)

    try:
        ex   = cdw_bridge.extract_from_cdw(cdw_path)
        result["_raw_text"] = getattr(ex, "pdf_text", "") or ""
        card = getattr(ex, "cdw_card", {}) or {}

        result["drawing_no"]    = card.get("marking", "")
        result["part_name"]     = card.get("name", "") or os.path.splitext(name)[0]
        result["material"]      = card.get("material", "")
        result["clean_mass_kg"] = _fmt_mass_ru(card.get("mass"))

        lib = _sortament_from_library(card)
        if lib:
            result["drawing_sortament"] = f"{lib} / {card.get('material', '')}".strip(" /")
        else:
            result["drawing_sortament"] = card.get("material", "")

        # 1) База эталонов.
        # Для CDW обозначение известно ТОЧНО — это свойство marking из файла,
        # поэтому ищем сразу по нему. Разбирать номер из текста не нужно и
        # опасно: _extract_drawing_no() рассчитан на формат 23.14.02.004 и
        # из «290426.00.00.141» выдаёт обрезанное «00.00.141».
        etalon = None
        marking = (card.get("marking") or "").strip()
        if marking:
            finder = getattr(drawing_db, "lookup_by_drawing_no", None)
            if finder:
                try:
                    etalon = finder(marking)
                except Exception as e:                       # noqa: BLE001
                    logger.warn(f"Поиск эталона по обозначению не удался: {e}")
        if not etalon:
            etalon = lookup_etalon(ex.pdf_text, ex.part_name)
        if etalon:
            db_name = etalon.get("part_name", "")
            if db_name and not result["part_name"]:
                result["part_name"] = db_name
            new_sortament, stock_size, stock_mass = parse_result_line(
                etalon.get("result_line", "")
            )
            result["new_sortament"]  = new_sortament
            result["stock_size"]     = stock_size
            result["stock_mass_kg"]  = stock_mass
            result["zagotovka_full"] = new_sortament
            result["source"]         = "CDW+база"
            return result

        # 2) Расчёт
        mark = normalize_mark(ex.material_mark or "")
        gost = normalize_gost(ex.material_gost or "")

        if not mark:
            result["status"] = "⚠ нет марки"
            result["error"]  = "Материал в штампе КОМПАСа пуст"
            result["source"] = "CDW"
            return result
        if not ex.length_mm or ex.length_mm <= 0:
            # Паспортные данные всё равно годные — отдаём их, отметив пробел
            result["status"] = "⚠ нет длины"
            result["error"]  = ("Размеры не параметризованы в чертеже — "
                                "длина заготовки не определена")
            result["source"] = "CDW"
            return result

        # Паспортные данные уже прочитаны и верны. Если расчёт заготовки
        # не вышел — это НЕ ошибка чтения файла, а нехватка исходных данных.
        # Строку сохраняем целиком, помечая только незакрытую часть.
        try:
            res = calculate_missing_stock_line(ex, ex.length_mm, mark, gost or "")
        except Exception as calc_err:
            result["status"] = "⚠ заготовка не определена"
            result["error"]  = str(calc_err)[:120]
            result["source"] = "CDW"
            logger.warn(f"CDW {name}: расчёт заготовки не удался: {calc_err}")
            return result

        new_sortament, stock_size, stock_mass = parse_result_line(res.result_line)

        result["new_sortament"]  = new_sortament
        result["stock_size"]     = stock_size
        result["stock_mass_kg"]  = stock_mass
        result["zagotovka_full"] = new_sortament
        result["source"]         = "CDW+расчёт"

    except Exception as e:
        result["status"] = "❌ ошибка"
        result["error"]  = str(e)[:120]
        logger.error(f"Пакетная обработка CDW: ошибка {name}: {e}")

    return result


# ============================================================
# Обработка одного .m3d/.a3d (КОМПАС, модель)
# ============================================================
#
# В .m3d геометрия закрыта (формат C3D) — ни длины, ни диаметра оттуда не
# получить без КОМПАСа/экспорта в STEP (см. ui/screens/model3d_tool.py и
# integrations/kompas/adapter.py). Здесь — то же самое честное поведение,
# что и в живом КОМПАС-адаптере: если деталь нашлась в базе эталонов по
# обозначению — отдаём полный результат как для CDW; если нет — отдаём то,
# что реально прочитано (материал, масса), и пишем прямым текстом, что
# заготовка не определена, а не подставляем случайное число.

def process_one_m3d(m3d_path: str) -> dict:
    from core.m3d_reader import read_m3d

    name = os.path.basename(m3d_path)
    result = _blank_result(m3d_path)

    try:
        info = read_m3d(m3d_path)
        if info is None:
            result["status"] = "❌ ошибка"
            result["error"]  = "Не удалось прочитать .m3d (не ZIP-контейнер КОМПАС, либо нет свойств детали)"
            return result

        result["_raw_text"]     = info.protocol
        result["drawing_no"]    = info.drawing_no
        result["part_name"]     = info.part_name or os.path.splitext(name)[0]
        result["material"]      = info.material
        result["clean_mass_kg"] = _fmt_mass_ru(info.mass_kg)

        etalon = None
        marking = (info.drawing_no or "").strip()
        if marking:
            finder = getattr(drawing_db, "lookup_by_drawing_no", None)
            if finder:
                try:
                    etalon = finder(marking)
                except Exception as e:                       # noqa: BLE001
                    logger.warn(f"Поиск эталона по обозначению не удался: {e}")

        if etalon:
            db_name = etalon.get("part_name", "")
            if db_name and not result["part_name"]:
                result["part_name"] = db_name
            new_sortament, stock_size, stock_mass = parse_result_line(
                etalon.get("result_line", "")
            )
            result["new_sortament"]  = new_sortament
            result["stock_size"]     = stock_size
            result["stock_mass_kg"]  = stock_mass
            result["zagotovka_full"] = new_sortament
            result["source"]         = "M3D+база"
            return result

        # Не нашлось в базе — геометрии в .m3d нет, заготовку определить
        # нечем. Это не ошибка чтения, а честная граница формата.
        result["status"] = "⚠ нет геометрии"
        result["error"]  = (
            "В .m3d нет геометрии (формат C3D закрыт) — заготовка не "
            "определена. Материал и масса прочитаны верно. Для полного "
            "расчёта нужен STEP этой же детали (вкладка «3D инструмент») "
            "или деталь должна быть в базе эталонов по обозначению."
        )
        result["source"] = "M3D"

    except Exception as e:
        result["status"] = "❌ ошибка"
        result["error"]  = str(e)[:120]
        logger.error(f"Пакетная обработка M3D: ошибка {name}: {e}")

    return result


# ============================================================
# ИИ-добор для сложных случаев
# ============================================================

def process_one_pdf_with_ai(pdf_path: str) -> dict:
    result = process_one_pdf(pdf_path)
    if result["status"] == "ОК":
        return result

    try:
        ex = extract_from_pdf(pdf_path, max_pages=3)
        pdf_short = ex.pdf_text[-800:].strip() if len(ex.pdf_text) > 800 else ex.pdf_text

        result["drawing_no"]        = result["drawing_no"] or extract_drawing_no(ex.pdf_text)
        result["drawing_sortament"] = result["drawing_sortament"] or extract_sortament_from_text(ex.pdf_text)
        result["clean_mass_kg"]     = result["clean_mass_kg"] or extract_clean_mass(ex.pdf_text)

        _apply_ai(result, ex, pdf_short, os.path.basename(pdf_path))

    except Exception as e:
        result["error"] = f"ИИ: {str(e)[:80]}"
        logger.error(f"Пакет ИИ ошибка {os.path.basename(pdf_path)}: {e}")

    return result


def process_one_cdw_with_ai(cdw_path: str) -> dict:
    result = process_one_cdw(cdw_path)
    if result["status"] == "ОК":
        return result

    try:
        ex = cdw_bridge.extract_from_cdw(cdw_path)
        _apply_ai(result, ex, ex.pdf_text[:1200], os.path.basename(cdw_path))
        if result["source"] == "ИИ":
            result["source"] = "CDW+ИИ"
    except Exception as e:
        result["error"] = f"ИИ: {str(e)[:80]}"
        logger.error(f"Пакет ИИ (CDW) ошибка {os.path.basename(cdw_path)}: {e}")

    return result


def _apply_ai(result: dict, ex, text_block: str, fname: str) -> None:
    """Общий ИИ-шаг для PDF и CDW. Меняет result на месте."""
    db_examples = build_etalon_examples_block("", "")
    gost_ctx    = get_gost_context_for_query(ex.pdf_text)

    prompt = (
        "Ты технолог ОТК. Определи недостающую заготовку для чертежа.\n\n"
        "ОТВЕТ строго 5 строк:\n"
        "Заготовка: [Тип Размер ГОСТ]\n"
        "Материал: [марка ГОСТ]\n"
        "Размер: [L=XXX или □XXX или BхL]\n"
        "Масса заготовки: [XX,XX кг]\n"
        "Сортамент: [Тип Размер ГОСТ]\n\n"
        f"{db_examples}\n{gost_ctx}\n"
        f"ЧЕРТЁЖ ({fname}):\n{text_block}\n\n"
        "Ответь строго по формату, только 5 строк."
    )

    answer = ai_ask(prompt, model="gpt-4o-mini")
    zagotovka = material = size = mass = sortament = ""

    for line in answer.strip().splitlines():
        line = line.strip()
        low  = line.lower()
        if   low.startswith("заготовка:"):  zagotovka = line.split(":", 1)[1].strip()
        elif low.startswith("материал:"):   material  = line.split(":", 1)[1].strip()
        elif low.startswith("размер:"):     size      = line.split(":", 1)[1].strip()
        elif low.startswith("масса заготовки:") or low.startswith("масса:"):
                                            mass      = line.split(":", 1)[1].strip()
        elif low.startswith("сортамент:"):  sortament = line.split(":", 1)[1].strip()

    if zagotovka:
        result["new_sortament"]  = sortament or zagotovka
        result["stock_size"]     = size
        result["stock_mass_kg"]  = re.sub(r"[^\d,\.]", "", mass)
        result["zagotovka_full"] = zagotovka
    else:
        result["zagotovka_full"] = answer.strip()[:120]

    result["source"] = "ИИ"
    result["status"] = "ОК"
    result["error"]  = ""
    logger.info(f"ИИ: {fname} → {result['zagotovka_full'][:50]}")


# ============================================================
# Диспетчер
# ============================================================

def _source_kind(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".cdw":
        return "cdw"
    if ext in (".m3d", ".a3d"):
        return "m3d"
    return "pdf"


def process_one_file(path: str, use_ai: bool = False) -> dict:
    """Единая точка входа: .cdw читаем структурно, .m3d — свойства без
    геометрии, остальное — через PDF."""
    kind = _source_kind(path)
    if kind == "cdw":
        result = process_one_cdw_with_ai(path) if use_ai else process_one_cdw(path)
    elif kind == "m3d":
        # У .m3d нет ИИ-добора: там нет ни текста чертежа, ни картинки,
        # чтобы модели было на чём угадывать — только структурные свойства.
        result = process_one_m3d(path)
    else:
        result = process_one_pdf_with_ai(path) if use_ai else process_one_pdf(path)

    raw_text = result.pop("_raw_text", "")
    activity_log.log_processing(
        source_path=path,
        mode="batch",
        status=result.get("status", ""),
        source_kind=kind,
        message=result.get("error", ""),
        used_ai=use_ai,
        result=result,
        raw_text=raw_text,
    )
    return result


SUPPORTED_PATTERNS = ("*.cdw", "*.CDW", "*.pdf", "*.PDF", "*.m3d", "*.M3D", "*.a3d", "*.A3D")

# Приоритет, если на одну деталь в папке лежит несколько форматов:
# .cdw читается точно и структурно (обозначение/материал/масса из файла,
# плюс размеры чертежа) — предпочитаем его всегда. .pdf — следующий
# (распознаётся текстом/ИИ, но хотя бы может нести готовый штамп).
# .m3d — последний: даёт только материал и массу, геометрии там нет.
_FORMAT_PRIORITY = {".cdw": 0, ".pdf": 1, ".m3d": 2, ".a3d": 2}


def scan_folder(folder: str) -> list:
    """
    Все чертежи/модели в папке. Если на одну деталь есть несколько
    форматов — берём наиболее точный (см. _FORMAT_PRIORITY).
    """
    from pathlib import Path
    found = []
    for pat in SUPPORTED_PATTERNS:
        found.extend(Path(folder).glob(pat))

    # Пропускаем служебные файлы-спутники macOS (AppleDouble, "._имя.cdw"),
    # которые появляются при копировании с Mac на флешку/сетевой диск —
    # это не чертежи, а метаданные, и чтение их как ZIP всегда падает.
    found = [f for f in found if not f.name.startswith("._")]

    by_stem = {}
    for f in sorted(found):
        prev = by_stem.get(f.stem)
        if prev is None or _FORMAT_PRIORITY.get(f.suffix.lower(), 9) < \
                _FORMAT_PRIORITY.get(prev.suffix.lower(), 9):
            by_stem[f.stem] = f
    return [str(p) for p in sorted(by_stem.values())]


# ============================================================
# Экспорт в Excel
# ============================================================

def export_to_excel(results: list, save_path: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise ImportError("openpyxl не установлен. pip install openpyxl")

    # Сортировка строк по номеру чертежа (как в таблице заказчика)
    results = sorted(results, key=lambda r: _drawing_sort_key(r.get("drawing_no", "")))

    wb = Workbook()
    ws = wb.active
    ws.title = "Заготовки"

    # ── Стили ────────────────────────────────────────────────
    thin = Border(
        left=Side(style="thin", color="808080"),
        right=Side(style="thin", color="808080"),
        top=Side(style="thin", color="808080"),
        bottom=Side(style="thin", color="808080"),
    )

    hdr1_font  = Font(name="Times New Roman", bold=True, size=11, color="FFFFFF")
    hdr2_font  = Font(name="Times New Roman", bold=True, size=10, color="FFFFFF")
    body_font  = Font(name="Times New Roman", size=10)
    title_font = Font(name="Times New Roman", bold=True, size=12)

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    hdr1_fill  = PatternFill("solid", fgColor="1E3A5F")   # тёмно-синий — группы
    hdr2_fill  = PatternFill("solid", fgColor="2D6A9F")   # синий — подзаголовки
    white      = PatternFill("solid", fgColor="FFFFFF")
    db_fill    = PatternFill("solid", fgColor="E3F2FD")
    ok_fill    = PatternFill("solid", fgColor="E8F5E9")
    ai_fill    = PatternFill("solid", fgColor="F3E8FF")
    cdw_fill   = PatternFill("solid", fgColor="E0F2F1")
    warn_fill  = PatternFill("solid", fgColor="FFF9C4")
    err_fill   = PatternFill("solid", fgColor="FFEBEE")

    # ── Заголовок документа ──────────────────────────────────
    ws.merge_cells("A1:H1")
    ws["A1"] = (f"OTK.AI — Ведомость заготовок   |   Дата: "
                f"{datetime.datetime.now().strftime('%d.%m.%Y %H:%M')}")
    ws["A1"].font      = title_font
    ws["A1"].alignment = center
    ws["A1"].fill      = white
    ws.row_dimensions[1].height = 24

    # ── Строка 2: группы столбцов ────────────────────────────
    ws.merge_cells("A2:A3")   # №
    ws.merge_cells("B2:C2")   # Чертёж
    ws.merge_cells("D2:G2")   # Заготовка
    ws.merge_cells("H2:H3")   # Источник

    for cell_ref, val in [("A2", "№"), ("B2", "Чертёж"),
                          ("D2", "Заготовка"), ("H2", "Источник\nданных")]:
        c = ws[cell_ref]
        c.value     = val
        c.font      = hdr1_font
        c.fill      = hdr1_fill
        c.alignment = center
        c.border    = thin

    for col in ["C2", "E2", "F2", "G2", "A3", "H3"]:
        ws[col].fill   = hdr1_fill
        ws[col].border = thin

    ws.row_dimensions[2].height = 22

    # ── Строка 3: подзаголовки ───────────────────────────────
    sub_headers = [
        ("B3", "№ чертежа"),
        ("C3", "Наименование"),
        ("D3", "Сортамент"),
        ("E3", "Размер\nзаготовки"),
        ("F3", "Масса\nзаг., кг"),
        ("G3", "Масса\nчист., кг"),
    ]
    for cell_ref, val in sub_headers:
        c = ws[cell_ref]
        c.value     = val
        c.font      = hdr2_font
        c.fill      = hdr2_fill
        c.alignment = center
        c.border    = thin

    ws.row_dimensions[3].height = 32

    # ── Данные ───────────────────────────────────────────────
    for i, r in enumerate(results, 1):
        row = i + 3
        src = r.get("source", "")

        if r["status"] == "ОК":
            if src.startswith("CDW"):
                fill = cdw_fill
            elif src == "расчёт":
                fill = ok_fill
            elif src.startswith("ИИ"):
                fill = ai_fill
            else:
                fill = db_fill
        elif "⚠" in r["status"]:
            fill = warn_fill
        else:
            fill = err_fill

        sortament = r.get("new_sortament", "") if r["status"] == "ОК" else r.get("error", "")

        values = {
            "A": i,
            "B": r.get("drawing_no", ""),
            "C": r.get("part_name", ""),
            "D": sortament,
            "E": r.get("stock_size", ""),
            "F": r.get("stock_mass_kg", ""),
            "G": r.get("clean_mass_kg", ""),
            "H": src if r["status"] == "ОК" else r["status"],
        }

        for col, val in values.items():
            cell = ws[f"{col}{row}"]
            cell.value     = val
            cell.font      = body_font
            cell.fill      = fill
            cell.border    = thin
            cell.alignment = left if col in ("B", "C", "D", "E") else center

        ws.row_dimensions[row].height = 22

    # ── Ширины столбцов ──────────────────────────────────────
    for col, w in {"A": 5, "B": 18, "C": 18, "D": 32,
                   "E": 18, "F": 12, "G": 12, "H": 14}.items():
        ws.column_dimensions[col].width = w

    ws.freeze_panes = "A4"

    # ── Легенда ──────────────────────────────────────────────
    lr = len(results) + 5
    ws.cell(row=lr, column=1, value="Легенда:").font = Font(name="Arial", bold=True)
    for j, (f_obj, label) in enumerate([
        (cdw_fill,  "Прочитано из файла КОМПАС (.cdw) — точные данные"),
        (db_fill,   "Из базы эталонов"),
        (ok_fill,   "Из расчёта по ГОСТ"),
        (ai_fill,   "Определено ИИ"),
        (warn_fill, "Предупреждение — часть данных не определена"),
        (err_fill,  "Ошибка обработки"),
    ]):
        r2 = lr + 1 + j
        ws.cell(row=r2, column=1).fill   = f_obj
        ws.cell(row=r2, column=1).border = thin
        ws.cell(row=r2, column=2, value=label).font = Font(name="Arial", size=10)

    # ── Лист «Легенда столбцов» — заказчик просил видеть все параметры ──
    ws2 = wb.create_sheet("Легенда столбцов")
    ws2["A1"] = "Откуда берётся каждый столбец"
    ws2["A1"].font = Font(name="Arial", bold=True, size=13)

    for col, name in zip("ABCD", ["Столбец", "Что означает", "Источник", "Надёжность"]):
        c = ws2[f"{col}3"]
        c.value     = name
        c.font      = hdr2_font
        c.fill      = hdr2_fill
        c.alignment = center
        c.border    = thin

    legend_rows = [
        ("№ чертежа", "Обозначение детали",
         "CDW: свойство marking. PDF: распознавание штампа",
         "CDW — точно, PDF — распознавание"),
        ("Наименование", "Наименование детали",
         "CDW: свойство name. PDF: распознавание штампа",
         "CDW — точно, PDF — распознавание"),
        ("Сортамент", "Вид проката, размер профиля и ГОСТ",
         "База эталонов → расчёт по ГОСТ → ИИ",
         "По базе — проверено, иначе требует проверки"),
        ("Размер заготовки", "Габарит заготовки под отрезку",
         "Расчёт по длине детали + припуск",
         "Требует проверки технолога"),
        ("Масса заг., кг", "Масса заготовки до обработки",
         "Расчёт: объём заготовки × плотность материала",
         "Требует проверки технолога"),
        ("Масса чист., кг", "Масса готовой детали",
         "CDW: свойство mass. PDF: поле «Масса» штампа",
         "CDW — точно, из файла"),
        ("Источник данных", "Каким слоем получена строка",
         "CDW / база / расчёт / ИИ",
         "Служебный столбец"),
    ]
    rr = 4
    for a, b, c, d in legend_rows:
        for col, val in zip("ABCD", (a, b, c, d)):
            cell = ws2[f"{col}{rr}"]
            cell.value     = val
            cell.font      = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border    = thin
        rr += 1

    for col, w in {"A": 22, "B": 36, "C": 48, "D": 32}.items():
        ws2.column_dimensions[col].width = w

    wb.save(save_path)


def write_run_log(results: list, log_path: str, folder_in: str = "") -> None:
    """Текстовый протокол рядом с таблицей: что обработано, что нет."""
    ok  = [r for r in results if r["status"] == "ОК"]
    bad = [r for r in results if r["status"] != "ОК"]
    lines = [
        "OTK.AI — протокол пакетной обработки",
        f"Дата:   {datetime.datetime.now().strftime('%d.%m.%Y %H:%M:%S')}",
        f"Папка:  {folder_in or '(список файлов)'}",
        f"Всего:  {len(results)}   Успешно: {len(ok)}   С замечаниями: {len(bad)}",
        "",
        "ФАЙЛЫ:",
    ]
    for r in results:
        lines.append(
            f"  [{r['status']}] {r['file']}"
            f"  | {r.get('drawing_no','')}"
            f"  | источник: {r.get('source','-')}"
            + (f"  | {r['error']}" if r.get("error") else "")
        )
    with open(log_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))


# ============================================================
# ЭКРАН
# ============================================================

class BatchToolScreen(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color="transparent")
        self.grid_rowconfigure(3, weight=1)
        self.grid_columnconfigure(0, weight=1)

        self._files: list[str] = []
        self._results: list[dict] = []
        self._processing = False

        self._folder_in   = ctk.StringVar(value="")
        self._folder_out  = ctk.StringVar(value="")

        self._build_ui()
        self._load_folders()

    # ---------- построение интерфейса ----------

    def _build_ui(self):
        top = ctk.CTkFrame(self, corner_radius=12,
                           fg_color=("gray92", "#141416"),
                           border_width=1, border_color=("gray80", "#222226"))
        top.grid(row=0, column=0, sticky="ew", padx=12, pady=(12, 6))
        top.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(top, text="Пакетная обработка чертежей",
                     font=ctk.CTkFont(size=16, weight="bold"),
                     text_color=("gray10", "gray90")
                     ).grid(row=0, column=0, padx=16, pady=(14, 4), sticky="w")

        ctk.CTkLabel(top,
                     text="Поддерживаются .cdw (КОМПАС), .pdf и .m3d/.a3d (КОМПАС, модель). "
                          "Из .cdw и .m3d обозначение, материал и масса читаются точно, без "
                          "распознавания. У .m3d нет геометрии — заготовка определится, только "
                          "если деталь уже есть в базе эталонов.",
                     font=ctk.CTkFont(size=12), text_color=("gray40", "#888888")
                     ).grid(row=1, column=0, columnspan=3, padx=16, pady=(0, 14), sticky="w")

        btn_frame = ctk.CTkFrame(top, fg_color="transparent")
        btn_frame.grid(row=0, column=2, padx=12, pady=12, sticky="e", rowspan=2)

        self.btn_add = ctk.CTkButton(
            btn_frame, text="Добавить файлы", height=34, corner_radius=10, width=140,
            fg_color=("gray88", "gray25"), hover_color=("gray80", "gray30"),
            text_color=("gray10", "gray90"), command=self.on_add_files)
        self.btn_add.pack(side="left", padx=4)

        self.btn_clear_files = ctk.CTkButton(
            btn_frame, text="Очистить список", height=34, corner_radius=10, width=130,
            fg_color=("gray88", "gray25"), hover_color=("gray80", "gray30"),
            text_color=("gray10", "gray90"), command=self.on_clear_files)
        self.btn_clear_files.pack(side="left", padx=4)

        self.btn_run = ctk.CTkButton(
            btn_frame, text="Запустить обработку", height=34, corner_radius=10, width=160,
            fg_color="#16a34a", hover_color="#15803d", text_color="white",
            command=self.on_run, state="disabled")
        self.btn_run.pack(side="left", padx=4)

        self.btn_export = ctk.CTkButton(
            btn_frame, text="Сохранить в Excel", height=34, corner_radius=10, width=150,
            fg_color=("gray88", "gray25"), hover_color=("gray80", "gray30"),
            text_color=("gray10", "gray90"), command=self.on_export, state="disabled")
        self.btn_export.pack(side="left", padx=4)

        self._use_ai = ctk.BooleanVar(value=False)
        ai_frame = ctk.CTkFrame(btn_frame, fg_color="transparent")
        ai_frame.pack(side="left", padx=(12, 4))
        ctk.CTkSwitch(ai_frame, text="ИИ для новых", variable=self._use_ai,
                      onvalue=True, offvalue=False, height=26,
                      font=ctk.CTkFont(size=12), text_color=("gray10", "gray90"),
                      progress_color="#16a34a").pack(side="left")
        ctk.CTkLabel(ai_frame, text="(медленнее)", font=ctk.CTkFont(size=10),
                     text_color=("gray50", "gray60")).pack(side="left", padx=(4, 0))

        # ---------- Режим папок ----------
        ff = ctk.CTkFrame(self, corner_radius=12,
                          fg_color=("gray92", "#141416"),
                          border_width=1, border_color=("gray80", "#222226"))
        ff.grid(row=1, column=0, sticky="ew", padx=12, pady=(0, 6))
        ff.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(ff, text="Обработка папки",
                     font=ctk.CTkFont(size=14, weight="bold"),
                     text_color=("gray10", "gray90")
                     ).grid(row=0, column=0, columnspan=3, padx=16, pady=(12, 2), sticky="w")

        ctk.CTkLabel(ff,
                     text="Читаем чертежи из Папки 1, сохраняем таблицу и протокол в Папку 2. "
                          "Оригиналы остаются без изменений.",
                     font=ctk.CTkFont(size=11), text_color=("gray40", "#888888")
                     ).grid(row=1, column=0, columnspan=3, padx=16, pady=(0, 8), sticky="w")

        rows = [
            ("Папка 1 — чертежи на вход",  self._folder_in),
            ("Папка 2 — куда таблицу",     self._folder_out),
        ]
        for i, (label, var) in enumerate(rows):
            r = i + 2
            ctk.CTkLabel(ff, text=label, font=ctk.CTkFont(size=12),
                         text_color=("gray30", "gray70"), width=190, anchor="w"
                         ).grid(row=r, column=0, padx=(16, 6), pady=3, sticky="w")
            ctk.CTkEntry(ff, textvariable=var, height=30, corner_radius=8
                         ).grid(row=r, column=1, padx=6, pady=3, sticky="ew")
            ctk.CTkButton(ff, text="Обзор", width=80, height=30, corner_radius=8,
                          fg_color=("gray88", "gray25"), hover_color=("gray80", "gray30"),
                          text_color=("gray10", "gray90"),
                          command=lambda v=var: self._pick_folder(v)
                          ).grid(row=r, column=2, padx=(6, 16), pady=3)

        self.btn_folder_run = ctk.CTkButton(
            ff, text="Обработать папку", height=34, corner_radius=10, width=180,
            fg_color="#2563eb", hover_color="#1d4ed8", text_color="white",
            command=self.on_run_folder)
        self.btn_folder_run.grid(row=4, column=0, columnspan=3, padx=16, pady=(8, 14), sticky="w")

        # ---------- Прогресс ----------
        pf = ctk.CTkFrame(self, corner_radius=12,
                          fg_color=("gray92", "#141416"),
                          border_width=1, border_color=("gray80", "#222226"))
        pf.grid(row=2, column=0, sticky="ew", padx=12, pady=(0, 6))
        pf.grid_columnconfigure(0, weight=1)

        self.lbl_status = ctk.CTkLabel(pf, text="Файлы не выбраны",
                                       font=ctk.CTkFont(size=12),
                                       text_color=("gray40", "#888888"), anchor="w")
        self.lbl_status.grid(row=0, column=0, padx=16, pady=(10, 4), sticky="w")
        self.progress = ctk.CTkProgressBar(pf, corner_radius=4, height=6,
                                           progress_color="#16a34a")
        self.progress.set(0)
        self.progress.grid(row=1, column=0, padx=16, pady=(0, 10), sticky="ew")

        # ---------- Таблица ----------
        tf = ctk.CTkFrame(self, corner_radius=12,
                          fg_color=("gray92", "#141416"),
                          border_width=1, border_color=("gray80", "#222226"))
        tf.grid(row=3, column=0, sticky="nsew", padx=12, pady=(0, 12))
        tf.grid_rowconfigure(1, weight=1)
        tf.grid_columnconfigure(0, weight=1)

        hdr = ctk.CTkFrame(tf, fg_color=("gray85", "#1a1a1f"), corner_radius=8)
        hdr.grid(row=0, column=0, sticky="ew", padx=10, pady=(10, 4))

        self._col_cfg = [
            ("№",                 4),
            ("№ чертежа",        16),
            ("Наименование",     16),
            ("Заготовка",        30),
            ("Новый сортамент",  20),
            ("Материал/штамп",   20),
            ("Размер",           14),
            ("Масса заг.",        9),
            ("Масса чист.",       9),
            ("Источник",         12),
        ]
        for ci, (hname, _) in enumerate(self._col_cfg):
            ctk.CTkLabel(hdr, text=hname, font=ctk.CTkFont(size=11, weight="bold"),
                         text_color=("gray30", "gray70"), anchor="w"
                         ).grid(row=0, column=ci, padx=(10 if ci == 0 else 4, 4),
                                pady=6, sticky="w")
        for ci, (_, w) in enumerate(self._col_cfg):
            hdr.grid_columnconfigure(ci, minsize=w * 7)

        self.scroll = ctk.CTkScrollableFrame(tf, corner_radius=8,
                                             fg_color=("gray95", "#0f0f11"))
        self.scroll.grid(row=1, column=0, sticky="nsew", padx=10, pady=(0, 10))
        for ci, (_, w) in enumerate(self._col_cfg):
            self.scroll.grid_columnconfigure(ci, minsize=w * 7)

        self._show_empty_hint()

    # ---------- настройки папок ----------

    def _load_folders(self):
        try:
            from core.settings import load_settings
            s = load_settings()
            self._folder_in.set(s.get("folder_in", "") or "")
            self._folder_out.set(s.get("folder_out", "") or s.get("export_folder", "") or "")
        except Exception:
            pass

    def _save_folders(self):
        try:
            from core.settings import load_settings, save_settings
            s = load_settings()
            s["folder_in"]   = self._folder_in.get()
            s["folder_out"]  = self._folder_out.get()
            save_settings(s)
        except Exception as e:
            logger.warn(f"Не удалось сохранить пути папок: {e}")

    def _pick_folder(self, var):
        d = filedialog.askdirectory(title="Выберите папку", initialdir=var.get() or None)
        if d:
            var.set(d)
            self._save_folders()

    # ---------- список файлов ----------

    def _show_empty_hint(self):
        ctk.CTkLabel(self.scroll,
                     text="Добавьте чертежи (.cdw, .pdf, .m3d/.a3d) либо укажите Папку 1 и нажмите «Обработать папку»",
                     font=ctk.CTkFont(size=13), text_color=("gray60", "gray50")
                     ).grid(row=0, column=0, columnspan=10, pady=40)

    def _clear_scroll(self):
        for w in self.scroll.winfo_children():
            w.destroy()

    def on_add_files(self):
        paths = filedialog.askopenfilenames(
            title="Выберите чертежи",
            filetypes=[("Чертежи КОМПАС, PDF и модели КОМПАС", "*.cdw *.pdf *.m3d *.a3d"),
                       ("КОМПАС чертёж", "*.cdw"),
                       ("PDF", "*.pdf"),
                       ("КОМПАС модель", "*.m3d *.a3d")])
        if not paths:
            return
        added = 0
        for p in paths:
            if p not in self._files:
                self._files.append(p)
                added += 1
        self._update_file_status()
        logger.info(f"Пакет: добавлено {added} файлов, всего {len(self._files)}")

    def on_clear_files(self):
        self._files.clear()
        self._results.clear()
        self._clear_scroll()
        self._show_empty_hint()
        self.progress.set(0)
        self.btn_run.configure(state="disabled")
        self.btn_export.configure(state="disabled")
        self._update_file_status()

    def _update_file_status(self):
        n = len(self._files)
        if n == 0:
            self.lbl_status.configure(text="Файлы не выбраны")
            self.btn_run.configure(state="disabled")
        else:
            n_cdw = sum(1 for f in self._files if cdw_bridge.is_cdw(f))
            self.lbl_status.configure(
                text=f"Выбрано файлов: {n} (из них .cdw: {n_cdw}) — нажмите «Запустить обработку»")
            self.btn_run.configure(state="normal")

    # ---------- запуск: список файлов ----------

    def on_run(self):
        if not self._files or self._processing:
            return
        self._start_worker(list(self._files), folder_mode=False)

    # ---------- запуск: папка ----------

    def on_run_folder(self):
        if self._processing:
            return
        f_in   = self._folder_in.get().strip()
        f_out  = self._folder_out.get().strip()

        if not f_in or not os.path.isdir(f_in):
            messagebox.showwarning("Папка 1", "Укажите существующую папку с чертежами.")
            return
        if not f_out or not os.path.isdir(f_out):
            messagebox.showwarning("Папка 2", "Укажите существующую папку для таблицы.")
            return

        files = scan_folder(f_in)
        if not files:
            messagebox.showinfo("Пусто", f"В папке нет .cdw, .pdf или .m3d/.a3d файлов:\n{f_in}")
            return

        if not messagebox.askyesno(
                "Обработать папку",
                f"Найдено чертежей: {len(files)}\n\n"
                f"Таблица и протокол будут сохранены в:\n{f_out}\n\n"
                "Исходные .cdw, .pdf и .m3d останутся на своих местах.\n\nПродолжить?"):
            return

        self._save_folders()
        self._start_worker(files, folder_mode=True)

    # ---------- общий воркер ----------

    def _start_worker(self, files: list, folder_mode: bool):
        self._processing = True
        self._results = []
        self._clear_scroll()
        self.btn_run.configure(state="disabled", text="Обработка...")
        self.btn_add.configure(state="disabled")
        self.btn_folder_run.configure(state="disabled")
        self.btn_export.configure(state="disabled")
        self.progress.set(0)

        use_ai = self._use_ai.get()
        f_out  = self._folder_out.get().strip()
        f_in   = self._folder_in.get().strip()

        def worker():
            total = len(files)
            for i, path in enumerate(files):
                self.after(0, self.lbl_status.configure,
                           {"text": f"Обработка {i + 1}/{total}: {os.path.basename(path)}"})
                self.after(0, self.progress.set, i / total)
                self._results.append(process_one_file(path, use_ai=use_ai))

            saved_path = ""
            save_report = ""
            if folder_mode:
                try:
                    saved_path, save_report = self._finalize_folder(f_out, f_in)
                except Exception as e:
                    save_report = f"ОШИБКА сохранения: {e}"
                    logger.error(f"Режим папки: {e}")

            def finish():
                self._results.sort(key=lambda r: _drawing_sort_key(r.get("drawing_no", "")))
                self._clear_scroll()
                for idx, r in enumerate(self._results):
                    self._add_result_row(idx, r)
                self.progress.set(1.0)
                ok  = sum(1 for r in self._results if r["status"] == "ОК")
                err = len(self._results) - ok
                txt = f"Готово: {ok} успешно, {err} с замечаниями | Всего: {len(self._results)}"
                if save_report:
                    txt += f" | {save_report}"
                self.lbl_status.configure(text=txt)
                self.btn_run.configure(state="normal", text="Запустить обработку")
                self.btn_add.configure(state="normal")
                self.btn_folder_run.configure(state="normal")
                if ok > 0:
                    self.btn_export.configure(state="normal")
                self._processing = False
                logger.info(f"Пакетная обработка завершена: {ok}/{len(self._results)} ОК")
                if folder_mode and saved_path:
                    messagebox.showinfo(
                        "Папка обработана",
                        f"Таблица сохранена:\n{saved_path}\n\n{save_report}")

            self.after(0, finish)

        threading.Thread(target=worker, daemon=True).start()

    def _finalize_folder(self, f_out: str, f_in: str):
        """Сохраняет Excel и протокол. Исходные чертежи не изменяются."""
        stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        xlsx = os.path.join(f_out, f"OTK_заготовки_{stamp}.xlsx")
        log_path = os.path.join(f_out, f"OTK_протокол_{stamp}.txt")

        # Исключаем перезапись файлов отчёта при повторной обработке.
        suffix = 1
        while os.path.exists(xlsx) or os.path.exists(log_path):
            xlsx = os.path.join(f_out, f"OTK_заготовки_{stamp}_{suffix}.xlsx")
            log_path = os.path.join(f_out, f"OTK_протокол_{stamp}_{suffix}.txt")
            suffix += 1

        export_to_excel(self._results, xlsx)
        write_run_log(self._results, log_path, f_in)
        return xlsx, "Исходные чертежи сохранены без изменений"

    # ---------- строка таблицы ----------

    def _add_result_row(self, idx: int, r: dict):
        src = r.get("source", "")
        if r["status"] == "ОК":
            if src.startswith("CDW"):
                row_color = ("#e0f2f1", "#14302c"); status_color = ("#00695c", "#4db6ac")
            elif src == "расчёт":
                row_color = ("#e8f5e9", "#1e3a2e"); status_color = ("#2e7d32", "#4ade80")
            elif src.startswith("ИИ"):
                row_color = ("#f3e8ff", "#2a1a3d"); status_color = ("#7c3aed", "#c084fc")
            else:
                row_color = ("#e3f2fd", "#1a2a3d"); status_color = ("#1565c0", "#64b5f6")
            text_color = ("gray10", "gray90")
        elif "⚠" in r["status"]:
            row_color = ("#fff9c4", "#3a3210"); text_color = ("gray10", "gray90")
            status_color = ("#f57f17", "#ffd54f")
        else:
            row_color = ("#ffebee", "#3a1515"); text_color = ("gray10", "gray90")
            status_color = ("#c62828", "#ef9a9a")

        values = [
            str(idx + 1),
            r.get("drawing_no", ""),
            r.get("part_name", ""),
            r.get("zagotovka_full", "") if r["status"] == "ОК" else r.get("error", ""),
            r.get("new_sortament", ""),
            r.get("drawing_sortament", "") or r.get("material", ""),
            r.get("stock_size", ""),
            r.get("stock_mass_kg", ""),
            r.get("clean_mass_kg", ""),
            src if r["status"] == "ОК" else r["status"],
        ]

        bg = ctk.CTkFrame(self.scroll, fg_color=row_color, corner_radius=6)
        bg.grid(row=idx, column=0, columnspan=10, sticky="ew", padx=4, pady=2)

        for ci, (val, (_, w)) in enumerate(zip(values, self._col_cfg)):
            ctk.CTkLabel(bg, text=val, font=ctk.CTkFont(size=11),
                         text_color=status_color if ci == 9 else text_color,
                         anchor="w", wraplength=w * 6
                         ).grid(row=0, column=ci, padx=(10 if ci == 0 else 4, 4),
                                pady=5, sticky="w")

        for ci, (_, w) in enumerate(self._col_cfg):
            bg.grid_columnconfigure(ci, minsize=w * 7)

    # ---------- ручной экспорт ----------

    def on_export(self):
        if not self._results:
            messagebox.showwarning("Нет данных", "Сначала запустите обработку.")
            return
        now = datetime.datetime.now().strftime("%Y-%m-%d_%H-%M")
        default_name = f"OTK_заготовки_{now}.xlsx"

        default_dir = self._folder_out.get().strip() or None
        if not default_dir:
            try:
                from core.settings import load_settings
                default_dir = load_settings().get("export_folder", "") or None
            except Exception:
                default_dir = None

        path = filedialog.asksaveasfilename(
            title="Сохранить результаты", defaultextension=".xlsx",
            filetypes=[("Excel файл", "*.xlsx")],
            initialfile=default_name,
            initialdir=default_dir)
        if not path:
            return
        try:
            export_to_excel(self._results, path)
            messagebox.showinfo("Сохранено",
                                f"Файл сохранён:\n{path}\n\nЗаписей: {len(self._results)}")
            logger.info(f"Экспорт в Excel: {path}")
            if messagebox.askyesno("Открыть файл?", "Открыть Excel файл?"):
                import subprocess, sys
                if sys.platform == "darwin":
                    subprocess.call(["open", path])
                elif sys.platform == "win32":
                    os.startfile(path)
                else:
                    subprocess.call(["xdg-open", path])
        except Exception as e:
            messagebox.showerror("Ошибка экспорта", str(e))
            logger.error(f"Ошибка экспорта: {e}")
