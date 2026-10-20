# -*- coding: utf-8 -*-
"""
Импорт исправленного Excel технолога.

Каждая строка «деталь -> правильная заготовка» делает две вещи:
  1) пополняет обучающую выборку помощника (тип заготовки);
  2) если есть номер чертежа, записывается эталоном в базу, чтобы при
     следующей обработке этой детали программа выдала ответ технолога.

Колонки ищутся по заголовкам (формат клиента и старая выгрузка).
"""

from __future__ import annotations

import re
from typing import Any, Dict, List, Optional, Tuple

from core.assistant import store

_COLS = {
    "name":  ("наименование", "название", "деталь"),
    "no":    ("обозначение", "номер чертеж", "№ чертежа", "децимальн", "чертеж", "чертёж"),
    "mat":   ("материал",),
    "len":   ("длина",),
    "stock": ("сортамент", "заготовка"),
    "size":  ("размер",),
    "mass":  ("масса заг", "вес заг"),
    "clean": ("масса чист",),
    "err":   ("ошибка",),
}

# Строки, где программа лишь предложила значение, а технолог не снял пометку,
# в базу не попадают: непроверенное предположение не должно стать эталоном.
_UNCONFIRMED_MARKS = ("длина предложена", "предположение ии", "размер взят у аналога",
                      "размер листа по контуру")


def stock_type_from_text(text: str) -> Optional[str]:
    t = (text or "").lower().replace("ё", "е")
    if "шестигран" in t: return "Шестигранник"
    if "швеллер" in t:   return "Швеллер"
    if "двутавр" in t or "балка" in t: return "Двутавр"
    if "уголок" in t:    return "УголокНеравн" if "неравн" in t else "Уголок"
    if "полос" in t:     return "Полоса"
    if "квадрат" in t:   return "Квадрат"
    if "круг" in t:      return "Круг"
    if "профиль" in t and ("30245" in t or "8639" in t or "8645" in t):
        return "ТрубаПроф"
    if "труба" in t:
        return "ТрубаПроф" if ("проф" in t or "8639" in t or "8645" in t) else "Труба"
    if "лист" in t:
        return "ЛистХК" if ("19904" in t or "хк" in t) else "ЛистГК"
    return None


def _norm_header(h: Any) -> str:
    return re.sub(r"\s+", " ", str(h or "")).strip().lower()


# Колонки, которые технолог мог добавить справа своим блоком ответов: если
# заголовок встречается дважды (колонка программы и колонка технолога),
# берётся ПОСЛЕДНЯЯ, то есть исправленная.
_LAST_WINS = ("stock", "size", "mass")


def _find_cols(header: List[str]) -> Dict[str, int]:
    found: Dict[str, int] = {}
    for i, h in enumerate(header):
        hl = _norm_header(h)
        for key, subs in _COLS.items():
            if any(s in hl for s in subs):
                if key not in found or key in _LAST_WINS:
                    found[key] = i
                break
    return found


def _num(v) -> Optional[float]:
    if v in (None, ""):
        return None
    try:
        x = float(str(v).replace(",", ".").replace("\xa0", "").replace(" ", ""))
    except ValueError:
        return None
    return x if x > 0 else None


def _fmt_ru(x: float) -> str:
    return f"{x:.2f}".replace(".", ",")


def build_etalon(drawing_no: str, name: str, stock_cell: str, size_cell: str,
                 mass_stock, mass_clean, material_fallback: str = "") -> Optional[Dict[str, Any]]:
    """Эталон из строки таблицы технолога. None, если тип заготовки не распознан."""
    stock_type = stock_type_from_text(stock_cell)
    if not stock_type or not drawing_no:
        return None
    sort_part, _, mat_part = (stock_cell or "").partition("/")
    sort_part = sort_part.strip().rstrip(";")
    material = (mat_part or material_fallback or "").strip()
    gost_m = re.search(r"ГОСТ\s*[\d\-–]+", sort_part, re.I)
    gost_stock = re.sub(r"\s+", "", gost_m.group(0)) if gost_m else ""
    d_m = re.search(r"(?:круг|квадрат|шестигранник|лист|полоса|труба|уголок|профиль)\s+(\d+(?:[.,]\d+)?)",
                    sort_part, re.I)
    d_blank = int(float(d_m.group(1).replace(",", "."))) if d_m else 0
    size_txt = (size_cell or "").strip()
    num = r"(\d+(?:[.,]\d+)?)"
    sheet_m = re.search(r"[HhBbНнВв]\s*[xхXХ×]\s*L\s*=?\s*" + num + r"\s*[xхXХ×]\s*" + num, size_txt)
    l_m = re.search(r"L\s*=\s*" + num, size_txt)
    if sheet_m:                      # лист: HxL = 62x108, берём больший габарит
        l_blank = max(float(sheet_m.group(1).replace(",", ".")),
                      float(sheet_m.group(2).replace(",", ".")))
    elif l_m:
        l_blank = float(l_m.group(1).replace(",", "."))
    else:
        dims = [float(x.replace(",", ".")) for x in re.findall(num, size_txt.split(",")[-1])]
        l_blank = max(dims) if dims else 0.0
    mark_m = re.match(r"\s*(.+?)\s+(ГОСТ\s*[\d\-–]+|ТУ\s*[\d\.\-]+)", material, re.I)
    mat_mark = mark_m.group(1).strip() if mark_m else material
    mat_gost = re.sub(r"\s+", "", mark_m.group(2)) if mark_m else ""
    mb = _num(mass_stock)
    mc = _num(mass_clean)
    result_line = f"{sort_part};\n{size_txt};\nМасса заготовки, кг: {_fmt_ru(mb) if mb else ''}"
    return {
        "drawing_no": drawing_no.strip(),
        "part_name": (name or "").strip(),
        "stock_type": stock_type,
        "gost_stock": gost_stock,
        "d_blank_std_mm": d_blank,
        "l_part_mm": l_blank,          # длина детали в таблице технолога не указана
        "l_blank_mm": l_blank,
        "material_mark": mat_mark,
        "material_gost": mat_gost,
        "mass_kg": round(mb, 3) if mb else 0.0,
        "mass_clean_kg": round(mc, 3) if mc else None,
        "result_line": result_line,
        "source": "excel_import",
    }


def import_corrected_excel(path: str) -> Tuple[int, int, int]:
    """Возвращает (добавлено решений, пропущено строк, записано эталонов)."""
    from openpyxl import load_workbook
    from core import drawing_db
    wb = load_workbook(path, read_only=True, data_only=True)
    added = skipped = etalons = 0
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hdr_i, cols = None, {}
        for i, r in enumerate(rows[:15]):
            c = _find_cols([str(x) if x is not None else "" for x in r])
            if "name" in c and "stock" in c:
                hdr_i, cols = i, c
                break
        if hdr_i is None:
            continue
        for r in rows[hdr_i + 1:]:
            get = lambda k: (r[cols[k]] if k in cols and cols[k] < len(r) else None)
            stock_cell = str(get("stock") or "")
            if stock_cell.strip().lower() in ("убрать", "убрать."):
                skipped += 1            # технолог убрал позицию (сборка и т.п.)
                continue
            st = stock_type_from_text(stock_cell)
            name = str(get("name") or "").strip()
            if not st or not name:
                skipped += 1
                continue
            if str(get("err") or "").strip().lower().startswith(_UNCONFIRMED_MARKS):
                skipped += 1
                continue
            ln = _num(get("len"))
            drawing_no = str(get("no") or "").strip()
            store.add_decision(name, str(get("mat") or ""), ln, st,
                               drawing_no=drawing_no, source="excel_import")
            added += 1
            try:
                from core import agent
                agent.learn_from_row(drawing_no, {"sortament": stock_cell,
                                                  "size": str(get("size") or "")})
            except Exception:                                  # noqa: BLE001
                pass
            et = build_etalon(drawing_no, name, stock_cell, str(get("size") or ""),
                              get("mass"), get("clean"), str(get("mat") or ""))
            if et:
                try:
                    drawing_db.upsert_by_drawing_no(et)
                    etalons += 1
                except Exception:                              # noqa: BLE001
                    pass
    return added, skipped, etalons
