import sys, tempfile, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
from core.assistant import model, excel_import, store


def test_stock_type_from_text():
    f = excel_import.stock_type_from_text
    assert f("Круг 130 ГОСТ2590-2006") == "Круг"
    assert f("Лист Б-ПН-10 ГОСТ19903") == "ЛистГК"
    assert f("Швеллер 10У") == "Швеллер"
    assert f("что-то") is None


def test_train_and_suggest(tmp_path, monkeypatch):
    monkeypatch.setattr(model, "MODELS_DIR", tmp_path)
    rep = model.train()
    assert rep["ok"] and rep["holdout_acc"] is not None
    s = model.suggest("Втулка", "В20", 300)
    assert s.candidates and s.candidates[0].stock_type == "Труба"


def test_gate_keeps_better_model(tmp_path, monkeypatch):
    monkeypatch.setattr(model, "MODELS_DIR", tmp_path)
    model.train()
    meta = model.active_meta()
    meta["holdout_acc"] = 1.0
    (tmp_path / f"{meta['version']}.json").write_text(
        __import__("json").dumps(meta), encoding="utf-8")
    rep = model.train()
    assert rep["replaced"] is False


def test_engine_forced_type():
    from core import engine
    ex = engine.Extracted(pdf_text="", part_name="Крышка", looks_rotational=False,
                          material_mark=None, material_gost=None, stamp_mass_kg=None,
                          diameter_features=[], max_d_mm=None, max_d_tol=None,
                          square_features_mm=[], max_square_mm=None, ra_values=[], length_mm=None)
    engine.FORCED_STOCK_TYPE = "Швеллер"
    try:
        assert engine.detect_stock_type(ex) == "Швеллер"
    finally:
        engine.FORCED_STOCK_TYPE = None
    assert engine.detect_stock_type(ex) != "Швеллер"


def test_decision_replaces_same_drawing(tmp_path, monkeypatch):
    from core.assistant import store as st
    monkeypatch.setattr(st, "ASSISTANT_DIR", tmp_path)
    st.add_decision("Вал", "45", 100, "Круг", "X.1")
    st.add_decision("Вал", "45", 100, "Квадрат", "X.1")
    d = st.load_decisions()
    assert len(d) == 1 and d[0]["stock_type"] == "Квадрат"


def test_client_format_rows():
    from core import client_format as cf
    res = [
        {"status": "ОК", "drawing_no": "290426.00.00.002", "part_name": "Штырь направляющий",
         "new_sortament": "Шестигранник 24 ГОСТ2879-2006", "stock_size": "L=150",
         "stock_mass_kg": "0,60", "clean_mass_kg": "0,33", "material": "Ст2пс ГОСТ 535-2005"},
        {"status": "⚠ нет длины", "drawing_no": "290426.00.00.021", "part_name": "Полоса",
         "error": "длина не определена", "clean_mass_kg": "0,14"},
    ]
    rows = cf.rows_for_client(res)
    assert rows[0][3] == "Шестигранник 24 ГОСТ 2879-2006/ Ст2пс ГОСТ 535-2005"
    assert rows[0][4] == "Шестигранник 24, L=150"
    assert rows[0][5] == "0,6" and rows[0][6] == "0,33"
    assert rows[1][3] == "" and rows[1][7] == "длина не определена"


def test_mass_check():
    from core import mass_check as mc
    assert mc.check("0,6", "0,33") == (True, "")
    assert mc.check("0,036", "0,005")[0] is True          # реальный эталон, отношение 7,2
    assert mc.check("0,2", "0,33")[0] is False            # легче детали
    assert mc.check("5", "0,3")[0] is False               # в 16 раз тяжелее
    assert mc.check("", "0,3") == (True, "")              # нет массы, проверки нет


def test_mass_check_on_real_etalons():
    import json, pathlib
    from core import mass_check as mc
    db = json.loads((pathlib.Path(__file__).resolve().parents[1] / "seed" / "drawing_db.json").read_text(encoding="utf-8"))
    for r in db:
        if r.get("mass_clean_kg"):
            assert mc.check(r["mass_kg"], r["mass_clean_kg"])[0], r["part_name"]


def test_formulas_first_does_not_use_rag(monkeypatch):
    from core import engine
    seen = {}

    def core_stub(ex, l, m, g):
        seen["rag"] = engine.USE_RAG
        return "ok"
    monkeypatch.setattr(engine, "_calculate_stock_core", core_stub)
    monkeypatch.setattr(engine, "_formulas_first_enabled", lambda: True)
    assert engine.calculate_missing_stock_line(None, 1, "", "") == "ok"
    assert seen["rag"] is False and engine.USE_RAG is True   # флаг восстановлен


def test_formulas_fail_then_ai_layer(monkeypatch):
    from core import engine
    calls = []

    def core_stub(ex, l, m, g):
        calls.append(engine.USE_RAG)
        if not engine.USE_RAG:
            raise ValueError("формулы не смогли")
        return "ai-answer"
    monkeypatch.setattr(engine, "_calculate_stock_core", core_stub)
    monkeypatch.setattr(engine, "_formulas_first_enabled", lambda: True)
    monkeypatch.setattr(engine, "_RAG_AVAILABLE", True)
    assert engine.calculate_missing_stock_line(None, 1, "", "") == "ai-answer"
    assert calls == [False, True]


def test_cdw_dims_sections_and_proposal():
    from core import cdw_dims as D
    sec = D.parse_section("Круг 40 ГОСТ 2590-2006/ 45 ГОСТ 1050-2013")
    assert sec.kind == "Круг" and sec.sizes == (40.0,)
    assert D.parse_section("Труба 108x20 ГОСТ 8732-2025").sizes == (108.0, 20.0)
    assert D.parse_section("Уголок 40х40х5 ГОСТ 8509-93").area_mm2 == 75 * 5
    assert D.parse_section("Лист 16 ГОСТ 19903-2015").kind == "Лист"
    # Штырь Ø40: деталь 95 мм, масса 0,65 кг -> заготовка 105
    p = D.propose_length(sec, [10.0, 35.0, 95.0], 0.65, 7850.0)
    assert p.part_length_mm == 95.0 and p.blank_length_mm == 105.0
    # Размеров нет или лист: предложения нет
    assert D.propose_length(sec, [], 0.65).blank_length_mm is None
    sheet = D.parse_section("Лист 4 ГОСТ 19903-2015")
    assert D.propose_length(sheet, [100.0], 1.0).blank_length_mm is None
    # Радиус не принимается за длину; нереальная масса отбрасывается
    big = D.parse_section("Круг 220 ГОСТ 2590-2006")
    p2 = D.propose_length(big, [110.0, 20.0], 5.64, 7850.0)
    assert p2.part_length_mm != 110.0


def test_atomic_json_and_corrupt_quarantine(tmp_path):
    from core import fsutil
    p = tmp_path / "db.json"
    fsutil.write_json_atomic(p, [{"a": 1}])
    assert p.read_text(encoding="utf-8").startswith("[")
    assert not (tmp_path / "db.json.tmp").exists()
    p.write_text("{битый", encoding="utf-8")
    kept = fsutil.quarantine_corrupt(p)
    assert kept and "corrupt" in kept and open(kept, encoding="utf-8").read() == "{битый"


def test_import_excel_writes_etalon_and_decision(tmp_path, monkeypatch):
    from core import drawing_db
    from core.assistant import store as st
    from openpyxl import Workbook
    monkeypatch.setattr(st, "ASSISTANT_DIR", tmp_path / "a")
    monkeypatch.setattr(drawing_db, "DB_PATH", str(tmp_path / "db.json"))
    monkeypatch.setattr(drawing_db, "SEED_DB_PATH", tmp_path / "no_seed.json")
    wb = Workbook(); ws = wb.active
    ws.append(["№", "№ чертежа", "Наименование", "Заготовка", "Размер заготовки",
               "Масса заг., кг", "Масса чист., кг"])
    ws.append([1, "140826.00.00.001", "Штырь",
               "Шестигранник 55 ГОСТ 2879-2006/ 12Х18Н10Т ГОСТ 5949-2018",
               "Шестигранник 55, L=170", "3,5", "2,5"])
    ws.append([2, "140826.00.00.099", "Прочее", "", "", "", "1,0"])     # без ответа: пропуск
    p = tmp_path / "t.xlsx"; wb.save(p)
    added, skipped, etalons = excel_import.import_corrected_excel(str(p))
    assert (added, skipped, etalons) == (1, 1, 1)
    r = drawing_db.lookup_by_drawing_no("140826.00.00.001")
    assert r["stock_type"] == "Шестигранник" and r["l_blank_mm"] == 170.0
    assert r["mass_kg"] == 3.5 and r["mass_clean_kg"] == 2.5
    assert r["result_line"].splitlines()[0].rstrip(";") == "Шестигранник 55 ГОСТ 2879-2006"
    # повторный импорт заменяет запись, а не дублирует
    excel_import.import_corrected_excel(str(p))
    assert len([x for x in drawing_db.load_db() if x["drawing_no"] == "140826.00.00.001"]) == 1


def test_thread_bridge_runs_posted_calls():
    from ui import thread_bridge
    import threading
    out = []
    t = threading.Thread(target=lambda: thread_bridge.post(out.append, 7))
    t.start(); t.join()
    assert out == []                                  # из потока ничего не выполнилось
    assert thread_bridge.drain_for_tests() == 1 and out == [7]


def test_forced_type_is_thread_local():
    import threading
    from core import engine
    seen = {}
    engine.set_forced_stock_type("Швеллер")
    try:
        t = threading.Thread(target=lambda: seen.update(v=getattr(engine._forced_tls, "value", None)))
        t.start(); t.join()
        assert seen["v"] is None                      # другой поток тип не видит
        assert getattr(engine._forced_tls, "value") == "Швеллер"
    finally:
        engine.set_forced_stock_type(None)


def test_seed_merge_adds_new_records_only(tmp_path, monkeypatch):
    import json
    from core import drawing_db as db
    seed = tmp_path / "seed.json"
    user = tmp_path / "db.json"
    monkeypatch.setattr(db, "SEED_DB_PATH", seed)
    monkeypatch.setattr(db, "DB_PATH", str(user))
    monkeypatch.setattr(db, "_seed_checked", False)
    rec = lambda n, h: {"drawing_no": n, "pdf_text_hash": h, "part_name": n}
    seed.write_text(json.dumps([rec("A.1", "h1")]), encoding="utf-8")
    assert [r["drawing_no"] for r in db.load_db()] == ["A.1"]          # первый запуск
    # пользователь удалил A.1 и добавил свою запись; в обновлении seed пришла A.2
    user.write_text(json.dumps([rec("U.1", "hu")]), encoding="utf-8")
    seed.write_text(json.dumps([rec("A.1", "h1"), rec("A.2", "h2")]), encoding="utf-8")
    monkeypatch.setattr(db, "_seed_checked", False)
    nums = sorted(r["drawing_no"] for r in db.load_db())
    assert nums == ["A.2", "U.1"]                                      # A.1 не возвращается


def test_agent_autonomy_is_earned(tmp_path, monkeypatch):
    from core import agent
    monkeypatch.setattr(agent, "AGENT_DIR", tmp_path)
    assert agent.autonomy("formula_stamp") == "проверка"
    for _ in range(30):
        agent.record_outcome("formula_stamp", True)
    assert agent.autonomy("formula_stamp") == "проверка"   # 30/30: нижняя граница ~0.89
    for _ in range(50):
        agent.record_outcome("formula_stamp", True)
    assert agent.autonomy("formula_stamp") == "авто"
    for _ in range(5):
        agent.record_outcome("formula_stamp", False)
    assert agent.autonomy("formula_stamp") == "проверка"
    for _ in range(100):
        agent.record_outcome("ai_guess", True)
    assert agent.autonomy("ai_guess") == "проверка"      # ИИ никогда не авто


def test_agent_learns_from_corrected_row(tmp_path, monkeypatch):
    import json
    from core import agent
    monkeypatch.setattr(agent, "AGENT_DIR", tmp_path)
    (tmp_path / "log").mkdir()
    rec = {"drawing_no": "АБВ 1.2", "source": "CDW+расчёт",
           "new_sortament": "Круг 40", "stock_size": "L=100"}
    (tmp_path / "log" / "d.jsonl").write_text(json.dumps(rec, ensure_ascii=False) + "\n", encoding="utf-8")
    ok = agent.learn_from_row("абв1.2", {"sortament": "Круг 40", "size": "L=100"}, tmp_path / "log")
    assert ok is True
    bad = agent.learn_from_row("АБВ 1.2", {"sortament": "Круг 45", "size": "L=100"}, tmp_path / "log")
    assert bad is False
    assert agent.load_stats()["formula_stamp"] == {"n": 2, "correct": 1}


def test_export_stamp_sortament_without_size():
    from core import client_format as cf
    r = {"drawing_no": "1.1", "part_name": "Нож", "status": "⚠ нет длины", "source": "CDW",
         "new_sortament": "Лист 6 ГОСТ 19903-2015",
         "material": "Лист 6 ГОСТ 19903-2015;12Х18Н10Т ГОСТ 7350-77", "clean_mass_kg": "5,04",
         "error": "габариты не заданы"}
    row = cf.rows_for_client([r])[0]
    assert row[3] == "Лист 6 ГОСТ 19903-2015/ 12Х18Н10Т ГОСТ 7350-77"
    assert row[4] == "" and row[7] == "габариты не заданы"


def test_import_takes_technologist_block_when_headers_repeat(tmp_path, monkeypatch):
    """В таблице технолога колонки программы (D-E) и его ответы (K-L) названы одинаково:
    берётся правый блок. Строка «Убрать» пропускается."""
    from core import drawing_db
    from core.assistant import store as st
    from openpyxl import Workbook
    monkeypatch.setattr(st, "ASSISTANT_DIR", tmp_path / "a")
    monkeypatch.setattr(drawing_db, "DB_PATH", str(tmp_path / "db.json"))
    monkeypatch.setattr(drawing_db, "SEED_DB_PATH", tmp_path / "no_seed.json")
    wb = Workbook(); ws = wb.active
    ws.append(["№", "№ чертежа", "Наименование", "Сортамент", "Размер заготовки",
               "Масса заг., кг", "Масса чист., кг", "", "Сортамент", "Размеры Заг", "Вес Заг"])
    ws.append([1, "1.1", "Нож", "Лист 2000x1000 ГОСТ 19903-74", "BхL=2000х1000", "196,00", "5,04", "",
               "Лист 6 ГОСТ 19903-2015/ 12Х18Н10Т ГОСТ 7350-77", "HxL=140x765", 5.07654])
    ws.append([2, "1.2", "Корпус СБ", "Лист 1", "x", "1", "1", "", "Убрать", None, None])
    p = tmp_path / "t.xlsx"; wb.save(p)
    added, skipped, etalons = excel_import.import_corrected_excel(str(p))
    assert (added, skipped, etalons) == (1, 1, 1)
    r = drawing_db.lookup_by_drawing_no("1.1")
    assert r["result_line"].splitlines()[1].rstrip(";") == "HxL=140x765" and r["d_blank_std_mm"] == 6


def test_import_client_header_and_sheet_size(tmp_path, monkeypatch):
    """Шапка технолога: «Чертеж» вместо «№ чертежа»; размер листа «HxL = 62x108»."""
    from core import drawing_db
    from core.assistant import store as st
    from openpyxl import Workbook
    monkeypatch.setattr(st, "ASSISTANT_DIR", tmp_path / "a")
    monkeypatch.setattr(drawing_db, "DB_PATH", str(tmp_path / "db.json"))
    monkeypatch.setattr(drawing_db, "SEED_DB_PATH", tmp_path / "no_seed.json")
    wb = Workbook(); ws = wb.active
    ws.append([]); ws.append([])
    ws.append(["№", "Чертеж", "Наименование", "Сортамент", "Размер заготовки",
               "Масса заг.", "Масса чист.", "Источник"])
    ws.append([1, "201647.03.10.001", "Кронштейн", "Лист 5 ГОСТ 19903-2015/ Ст3пс5 ГОСТ 14637-89",
               "HxL = 62x108", 0.263, 0.255, "Технолог"])
    p = tmp_path / "t.xlsx"; wb.save(p)
    assert excel_import.import_corrected_excel(str(p)) == (1, 0, 1)
    r = drawing_db.lookup_by_drawing_no("201647.03.10.001")
    assert r["l_blank_mm"] == 108.0 and r["d_blank_std_mm"] == 5
    assert "62x108" in r["result_line"]


def test_analog_found_by_gabarit_not_by_name():
    from core import analogs
    def et(no, mc, size="HxL=865x1005", name="Стенка"):
        return {"drawing_no": no, "part_name": name, "mass_clean_kg": mc, "mass_kg": 27.3,
                "result_line": f"Лист 4 ГОСТ 19903-2015;\n{size};\nМасса заготовки, кг: 27,30"}
    db = [et("1", 26.0), et("2", 26.8, name="Совсем другое имя")]
    dims = (860.0, 1000.0)
    a = analogs.find_analog("Любое", "Лист 4 ГОСТ 19903-2015", 26.0, db, "9", part_dims=dims)
    assert a.n == 2 and a.size_text == "HxL=865x1005" and a.confidence == "высокая"
    # габарит детали не помещается в заготовку / заготовка слишком велика: аналога нет
    assert analogs.find_analog("Стенка", "Лист 4 ГОСТ 19903-2015", 26.0, db, "9", part_dims=(900.0, 1100.0)).n == 0
    assert analogs.find_analog("Стенка", "Лист 4 ГОСТ 19903-2015", 26.0, db, "9", part_dims=(400.0, 500.0)).n == 0
    assert analogs.find_analog("Стенка", "Лист 4 ГОСТ 19903-2015", 3.6, db, "9", part_dims=dims).n == 0   # масса другая
    assert analogs.find_analog("Стенка", "Лист 6 ГОСТ 19903-2015", 26.0, db, "9", part_dims=dims).n == 0  # другая толщина
    assert analogs.find_analog("Стенка", "Лист 4 ГОСТ 19903-2015", 26.0, db, "1", part_dims=dims).n == 1  # сам себе не аналог
    # без габаритов: допуск по массе вдвое жёстче, уверенность низкая
    b = analogs.find_analog("Стенка", "Лист 4 ГОСТ 19903-2015", 26.0, db[:1], "9")
    assert b.n == 1 and b.confidence == "низкая"


def test_allowance_is_learned_only_when_consistent(tmp_path, monkeypatch):
    from core import allowance
    monkeypatch.setattr(allowance, "AGENT_DIR", tmp_path)
    monkeypatch.setattr(allowance, "from_db", lambda db=None: {})
    assert allowance.get("Труба", 10) == (10.0, 0, False)
    for d in (10, 10, 10, 10):
        allowance.record("Труба", d)
    assert allowance.get("Труба", 10)[2] is False          # мало наблюдений
    allowance.record("Труба", 10)
    assert allowance.get("Труба", 0) == (10.0, 5, True)
    for d in (0, 0, 0, 0, 0):
        allowance.record("Круг", d)
    assert allowance.get("Круг", 10) == (0.0, 5, True)
    for d in (0, 30, 5, 40, 10):
        allowance.record("Профиль", d)
    assert allowance.get("Профиль", 10)[2] is False        # разброс большой: остаётся значение по умолчанию


def test_contour_sheet_size_and_mass_gate(monkeypatch):
    from core import cdw_contour as cc
    corners = [(-382.5, 70.0), (382.5, 70.0), (-382.5, -70.0), (382.5, -70.0), (-305.0, 26.0)]
    monkeypatch.setattr(cc, "_pairs", lambda p: corners)
    c = cc.propose_sheet("x.cdw", "Лист 6 ГОСТ 19903-2015", 5.04, 7.9)
    assert c and (c.h_mm, c.l_mm) == (140.0, 765.0) and c.method == "прямоугольник"
    # масса не сходится (деталь втрое легче габарита): предложения нет
    assert cc.propose_sheet("x.cdw", "Лист 6", 1.5, 7.9) is None
    assert cc.propose_sheet("x.cdw", "Круг 30", 5.04, 7.9) is None


def test_contour_is_never_auto():
    from core import agent
    assert agent.autonomy("contour") != "авто" if hasattr(agent, "autonomy") else True
    assert "contour" in agent.NEVER_AUTO


def test_client_xlsx_keeps_note_for_proposals():
    from core import client_format as cf
    base = {"drawing_no": "1", "part_name": "Нож", "new_sortament": "Лист 6 ГОСТ 19903-2015",
            "zagotovka_full": "Лист 6 ГОСТ 19903-2015", "stock_size": "HxL=140x765",
            "stock_mass_kg": "5,04", "clean_mass_kg": 5.04, "error": "Размер листа по контуру"}
    prop = cf.rows_for_client([dict(base, status="⚠ контур: проверить")])[0]
    ok = cf.rows_for_client([dict(base, status="ОК")])[0]
    assert "контуру" in prop[-1] and ok[-1] == ""


def test_allowance_prior_from_db(tmp_path, monkeypatch):
    from core import allowance
    monkeypatch.setattr(allowance, "_path", lambda: tmp_path / "a.json")
    db = [{"stock_type": "Труба", "l_part_mm": 100 + i, "l_blank_mm": 120 + i} for i in range(6)]
    db += [{"stock_type": "Круг", "l_part_mm": 50, "l_blank_mm": 70}] * 2      # мало
    db += [{"stock_type": "ЛистГК", "l_part_mm": 50, "l_blank_mm": 50}] * 8     # листы не считаются
    assert allowance.get("Труба", 10, db) == (20.0, 6, True)
    assert allowance.get("Круг", 10, db)[2] is False
    assert "ЛистГК" not in allowance.from_db(db)
    for _ in range(5):                              # исправления технолога важнее базы
        allowance.record("Труба", 10)
    assert allowance.get("Труба", 10, db)[0] == 10.0


def test_crosscheck_cdw_vs_db():
    from core import crosscheck as cx
    et = {"result_line": "Лист 4 ГОСТ 19903-2015;\nHxL=865x1005;", "mass_clean_kg": 26.0}
    assert cx.compare({"material": "Лист 4 ГОСТ 19903-2015;Ст3пс", "mass": "26"}, et) == []
    n = cx.compare({"material": "Лист 6 ГОСТ 19903-2015;Ст3пс", "mass": "40"}, et)
    assert len(n) == 2
    assert cx.missing_clean_mass({"mass": "5,5"}, {"result_line": ""}) == 5.5
    assert cx.missing_clean_mass({"mass": "5,5"}, et) is None


def test_trace_writes_and_exports(tmp_path, monkeypatch):
    from core import trace
    monkeypatch.setattr(trace, "PROCESSING_LOG_DIR", tmp_path)
    trace.begin("/x/деталь.cdw")
    trace.step("аналог", "не найден", n=0)
    out = tmp_path / "out.txt"
    assert trace.export_text(str(out)) > 2
    txt = out.read_text(encoding="utf-8")
    assert "деталь.cdw" in txt and "аналог" in txt


def test_analog_reason_and_contour_debug():
    from core import analogs
    a = analogs.find_analog("Стенка", "Лист 4 ГОСТ 19903-2015", 26.0, [], "1")
    assert not a.n and "подходящей заготовки нет" in a.reason


def test_build_version_reads_file_or_git(tmp_path, monkeypatch):
    from core import version
    f = tmp_path / "build_info.txt"
    f.write_text("abc1234  2026-10-21 12:00\n", encoding="utf-8")
    monkeypatch.setattr(version, "_candidates", lambda: iter([f]))
    assert version.get_build() == "abc1234  2026-10-21 12:00"
    monkeypatch.setattr(version, "_candidates", lambda: iter([tmp_path / "нет.txt"]))
    assert version.get_build()


def test_contour_roughness_allowance(monkeypatch):
    from core import cdw_contour as cc
    corners = [(-382.5, 70.0), (382.5, 70.0), (-382.5, -70.0), (382.5, -70.0)]
    monkeypatch.setattr(cc, "_pairs", lambda p: corners)
    base = ("x.cdw", "Лист 6 ГОСТ 19903-2015", 5.04, 7.9)
    assert cc.parse_ra("Ra 12,5") == 12.5 and cc.parse_ra("Rz 40") is None and cc.parse_ra("") is None
    c50 = cc.propose_sheet(*base, ra=50.0)
    assert (c50.h_mm, c50.l_mm) == (140.0, 765.0) and "без припуска" in c50.note
    c25 = cc.propose_sheet(*base, ra=25.0)
    assert (c25.h_mm, c25.l_mm) == (150.0, 775.0)
    cn = cc.propose_sheet(*base)
    assert (cn.h_mm, cn.l_mm) == (140.0, 765.0) and "не прочитана" in cn.note


def test_fill_marking_from_filename():
    import sys, types
    from unittest import mock
    for m in ("customtkinter", "tkinter", "tkinter.filedialog", "tkinter.messagebox"):
        sys.modules.setdefault(m, mock.MagicMock())
    from ui.screens import batch_tool as bt
    card, res = {}, {}
    bt._fill_from_filename(card, "/x/935028.01.08.005 Прокладка.cdw", res)
    assert card["marking"] == "935028.01.08.005" and card["name"] == "Прокладка" and "имени файла" in res["note"]
    card2, res2 = {"marking": "1.2.3", "name": "Деталь"}, {}
    bt._fill_from_filename(card2, "/x/файл.cdw", res2)
    assert not res2
